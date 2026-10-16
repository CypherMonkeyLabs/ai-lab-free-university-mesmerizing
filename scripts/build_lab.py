#!/usr/bin/env python3
"""Build the AI omni-channel learning-lab front door from scripts/rooms_data.py."""

from __future__ import annotations

import html
import json
import sys
import zipfile
from pathlib import Path

import cairosvg
from docx import Document
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from rooms_data import TRACKS  # noqa: E402
from source_catalog import (  # noqa: E402
    CHECKED,
    OVERVIEW_KEYS,
    SHEET_KEYS,
    TEMPLATE_KEYS,
    entries,
    go_deeper_html,
    room_keys,
    source_lines,
)

ROOT = Path(__file__).resolve().parents[1]
SITE = "https://cyphermonkeylabs.com"
CONTACT = "https://cyphermonkeylabs.com/#contact"
GITHUB = "https://github.com/CypherMonkeyLabs"
REPO = "https://github.com/CypherMonkeyLabs/ai-lab-free-university-mesmerizing"
LINKTREE = "https://linktr.ee/cypherMonkey"
DISC = (
    "This is an independent, free, learning-only student project. "
    "It respects every company and vendor, and every trademark belongs to its owner. "
    "It is not affiliated with or endorsed by any of them. "
    "If anyone has a concern, the content will be removed or corrected promptly via the contact page. "
    "Nothing here is sold. The author is happy to volunteer. "
    "Notes about recording, consent, and emergency calling are educational and are not legal advice."
)
COPY = "© 2026 CypherMonkeyLabs 🐒"
RECORDS: list[dict] = []


def esc(text: object) -> str:
    return html.escape(str(text), quote=True)


def write_text(rel: str, text: str, kind: str) -> None:
    path = ROOT / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    RECORDS.append({"rel": rel, "kind": kind, "bytes": path.stat().st_size, "disclaimer": "learning-only student project" in text})


def write_bytes(rel: str, data: bytes, kind: str, disclaimer: bool) -> None:
    path = ROOT / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    RECORDS.append({"rel": rel, "kind": kind, "bytes": len(data), "disclaimer": disclaimer})


def svg_png(rel_svg: str, svg: str) -> None:
    write_text(rel_svg, svg, "SVG")
    png_rel = str(Path(rel_svg).with_suffix(".png"))
    png = cairosvg.svg2png(bytestring=svg.encode("utf-8"), output_width=1200)
    write_bytes(png_rel, png, "PNG", True)


def svg_pdf(rel: str, svg: str) -> None:
    pdf = cairosvg.svg2pdf(bytestring=svg.encode("utf-8"))
    write_bytes(rel, pdf, "PDF", True)


def prefix(depth: int) -> str:
    return "../" * depth


def write_room_sources(rid: str, title: str, keys: list[str]) -> str:
    rel = f"rooms/sources/{rid}/sources.json"
    payload = {"room": rid, "title": title, "last_checked": CHECKED, "sources": entries(keys)}
    write_text(rel, json.dumps(payload, indent=2) + "\n", "JSON")
    return rel


def monkey_icon() -> str:
    return (
        '<svg class="monkey-icon" viewBox="0 0 32 32" aria-hidden="true">'
        '<circle cx="16" cy="18" r="8" fill="#fbbf24" stroke="#b45309" stroke-width="1.4"/>'
        '<circle cx="8" cy="12" r="4" fill="#f59e0b" stroke="#9a3412" stroke-width="1.2"/>'
        '<circle cx="24" cy="12" r="4" fill="#f59e0b" stroke="#9a3412" stroke-width="1.2"/>'
        '<circle cx="13" cy="17" r="1.3" fill="#1c1408"/>'
        '<circle cx="19" cy="17" r="1.3" fill="#1c1408"/>'
        '<circle cx="13" cy="17" r="2.6" fill="none" stroke="#0e7490" stroke-width="1.1"/>'
        '<circle cx="19" cy="17" r="2.6" fill="none" stroke="#0e7490" stroke-width="1.1"/>'
        '<path d="M15.4 17h1.2" stroke="#0e7490" stroke-width="1" fill="none"/>'
        '<path d="M13.2 21c1.2 1.2 4.4 1.2 5.6 0" fill="none" stroke="#1c1408" stroke-width="1.2" stroke-linecap="round"/>'
        "</svg>"
    )


def linktree_icon() -> str:
    return (
        '<svg viewBox="0 0 24 24" aria-hidden="true">'
        '<path fill="currentColor" d="M13.2 2.5h-2.4v5.1L7 3.9 5.3 5.6l3.9 3.9H3.6v2.4h5.7L5.3 15.9'
        'l1.7 1.7 4-4v7.9h2.4v-7.9l4 4 1.7-1.7-4-4h5.7V9.5h-5.6l3.9-3.9L17 3.9l-3.8 3.7z"/>'
        "</svg>"
    )


def github_icon() -> str:
    return (
        '<svg viewBox="0 0 16 16" aria-hidden="true"><path fill="currentColor" d="M8 0C3.58 0 0 3.58 0 8c0 3.54 2.29 6.53'
        " 5.47 7.59.4.07.55-.17.55-.38 0-.19-.01-.82-.01-1.49-2.01.37-2.53-.49-2.69-.94-.09-.23-.48-.94-.82-1.13-.28-.15"
        "-.68-.52-.01-.53.63-.01 1.08.58 1.23.82.72 1.21 1.87.87 2.33.66.07-.52.28-.87.51-1.07-1.78-.2-3.64-.89-3.64-3.95"
        " 0-.87.31-1.59.82-2.15-.08-.2-.36-1.02.08-2.12 0 0 .67-.21 2.2.82.64-.18 1.32-.27 2-.27.68 0 1.36.09 2 .27 1.53-1.04"
        ' 2.2-.82 2.2-.82.44 1.1.16 1.92.08 2.12.51.56.82 1.27.82 2.15 0 3.07-1.87 3.75-3.65 3.95.29.25.54.73.54 1.48'
        ' 0 1.07-.01 1.93-.01 2.2 0 .21.15.46.55.38A8.01 8.01 0 0 0 16 8c0-4.42-3.58-8-8-8z"/></svg>'
    )


def brand_row() -> str:
    return (
        f'<p class="footer-brand">CypherMonkey Labs</p>'
        f'<nav class="brand-icons" aria-label="CypherMonkey Labs">'
        f'<a class="brand-icon brand-web" href="{SITE}" aria-label="CypherMonkey Labs website">{monkey_icon()}</a>'
        f'<a class="brand-icon brand-gh" href="{GITHUB}" aria-label="CypherMonkey Labs on GitHub">{github_icon()}</a>'
        f'<a class="brand-icon brand-lt" href="{LINKTREE}" aria-label="Linktree">{linktree_icon()}</a>'
        f"</nav><p class=\"site-copy\">{esc(COPY)}</p>"
    )


def disclaimer_panel() -> str:
    return (
        '<aside class="disclaimer-panel" aria-labelledby="disc-title">'
        '<span class="disclaimer-ico" aria-hidden="true"><svg viewBox="0 0 24 24"><path d="M12 2.5 4.5 5.5v5.8c0 4.6 3.1 8.6 7.5 10.2 4.4-1.6 7.5-5.6 7.5-10.2V5.5z" fill="currentColor" opacity=".18" stroke="currentColor" stroke-width="1.6"/></svg></span>'
        '<div><h2 id="disc-title">A note from the student behind this lab</h2>'
        f"<p>{esc(DISC)}</p>"
        f'<p>Contact: <a href="{CONTACT}">cyphermonkeylabs.com/#contact</a>. '
        '<a href="https://creativecommons.org/licenses/by-nc-sa/4.0/">CC BY-NC-SA 4.0</a> content. MIT code.</p>'
        "</div></aside>"
    )


def footer() -> str:
    return f'<footer class="site-footer lab-footer" data-canon-footer>{disclaimer_panel()}{brand_row()}</footer>'


NAV = [
    ("home", "Home", "index.html", "The lifelong-learner front door."),
    ("learn", "Learn", "learn.html", f"{len(TRACKS)} tracks, from files through voice and the contact center."),
    ("voice", "Voice stack", "voice-stack.html", "STT, a model, speech, and the transport."),
    ("analytics", "Analytics", "speech-analytics.html", "Transcription, sentiment, QA scores, and redaction."),
    ("safety", "Safety", "agent-safety.html", "Handoff, guardrails, consent, and tests."),
    ("diagrams", "Diagrams", "diagrams.html", "Every drawing as SVG, PNG, and HTML."),
    ("flows", "Flows", "flows.html", "Animated paths you can export."),
    ("play", "Play", "play.html", "A quiz, a match, and a sort."),
    ("listen", "Listen", "listen.html", "Read-along that follows the page."),
    ("downloads", "Downloads", "downloads.html", "Sheets, cards, templates, and the zip."),
    ("ecosystem", "Ecosystem", "ecosystem.html", "Glossary, sisters, maps, and bridges."),
    ("glossary", "Glossary", "glossary.html", "Terms from the rooms."),
    ("certs", "Certs", "certs.html", "Confirm the fee. This site awards none."),
    ("labs", "Labs", "labs.html", "Zero-key Python in this repo."),
    ("help", "Help", "help.html", "Menu, course, guide, and the tour."),
    ("legal", "Legal", "legal.html", "Licence, trademarks, and the disclaimer."),
]


def icon(name: str) -> str:
    paths = {
        "home": "M4 11 L12 4 L20 11 V20 H4 Z",
        "learn": "M5 5 H19 V19 H5 Z M8 9 H16 M8 13 H13",
        "voice": "M12 4 A3 3 0 0 1 15 7 V12 A3 3 0 0 1 9 12 V7 A3 3 0 0 1 12 4 M8 12 A4 4 0 0 0 16 12 M12 16 V20 M9 20 H15",
        "analytics": "M4 19 H20 M7 16 V10 M12 16 V6 M17 16 V12",
        "safety": "M12 3 L19 6 V12 C19 16 16 19 12 21 C8 19 5 16 5 12 V6 Z M9 12 L11 14 L15 10",
        "diagrams": "M4 16 L10 8 L14 12 L20 6 M4 20 H20",
        "flows": "M4 7 H11 V11 H4 Z M13 13 H20 V17 H13 Z",
        "play": "M8 5 V19 L19 12 Z",
        "listen": "M5 10 V14 H8 L13 18 V6 L8 10 Z",
        "downloads": "M12 4 V14 M8 11 L12 15 L16 11 M5 19 H19",
        "ecosystem": "M12 20 V10 M12 10 C12 6 16 4 18 4 C17 8 14 10 12 10 Z",
        "glossary": "M6 4 H16 V20 H6 Z M6 4 C6 4 8 6 8 12",
        "certs": "M12 3 L14 8 L19 8 L15 12 L17 17 L12 14 L7 17 L9 12 L5 8 L10 8 Z",
        "labs": "M8 4 H16 L15 10 L19 18 H5 L9 10 Z",
        "help": "M12 17 H12.1 M9 9 A3 3 0 0 1 15 10 C15 12 12 12.5 12 14",
        "legal": "M12 3 L19 6 V12 C19 16 16 19 12 21 C8 19 5 16 5 12 V6 Z",
        "menu": "M4 7 H20 M4 12 H20 M4 17 H20",
    }
    return f'<svg viewBox="0 0 24 24" aria-hidden="true"><path d="{paths[name]}" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linejoin="round" stroke-linecap="round"/></svg>'


def chrome(title: str, main: str, depth: int, current: str, room: str = "") -> str:
    pre = prefix(depth)
    chips = []
    drawer = []
    for key, label, href, desc in NAV:
        here = ' aria-current="page"' if key == current else ""
        tour = f' data-tour="{key}"' if key in {"home", "learn", "diagrams", "play", "listen", "downloads"} else ""
        chips.append(
            f'<a class="section-chip" href="{pre}{href}"{here}{tour}><span class="chip-ico">{icon(key)}</span><span>{label}</span></a>'
        )
        cur = " is-current" if key == current else ""
        drawer.append(
            f'<li><a class="{cur.strip()}" href="{pre}{href}"{here}>'
            f'<span class="edge-ico">{icon(key)}</span>'
            f'<span class="edge-copy"><span class="edge-label">{label}</span>'
            f'<span class="edge-desc">{esc(desc)}</span></span></a></li>'
        )
    room_attr = f' data-room="{esc(room)}"' if room else ""
    dock = (
        '<nav class="phone-dock" aria-label="Phone">'
        f'<button type="button" class="dock-item" data-open-menu><span class="menu-ico">{icon("menu")}</span>Menu</button>'
        f'<a class="dock-item" href="{pre}learn.html" data-tour="learn"><span class="menu-ico">{icon("learn")}</span>Learn</a>'
        f'<a class="dock-item" href="{pre}diagrams.html" data-tour="diagrams"><span class="menu-ico">{icon("diagrams")}</span>Diagrams</a>'
        f'<a class="dock-item" href="{pre}play.html" data-tour="play"><span class="menu-ico">{icon("play")}</span>Play</a>'
        f'<a class="dock-item" href="{pre}listen.html" data-tour="listen"><span class="menu-ico">{icon("listen")}</span>Listen</a>'
        f'<button type="button" class="dock-item" data-open-guide><span class="menu-ico">{icon("help")}</span>Guide</button>'
        "</nav>"
    )
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="theme-color" content="#fff7ed">
<title>{esc(title)}</title>
<script>try{{var t=localStorage.getItem("ai-theme");if(["dark","contrast","holographic","cyberpunk","neon","sunset"].indexOf(t)<0)t="bright";document.documentElement.dataset.theme=t;var f=localStorage.getItem("ai-font");if(["sm","md","lg","xl"].indexOf(f)<0)f="md";document.documentElement.dataset.font=f;var m=localStorage.getItem("ai-motion");if(m==="reduce"||m==="full")document.documentElement.dataset.motion=m;var ink=localStorage.getItem("ai-ink");if(ink&&/^#[0-9a-f]{{6}}$/i.test(ink))document.documentElement.style.setProperty("--ink",ink);}}catch(e){{document.documentElement.dataset.theme="bright";document.documentElement.dataset.font="md";}}</script>
<link rel="icon" href="{pre}favicon.svg" type="image/svg+xml">
<link rel="manifest" href="{pre}manifest.webmanifest">
<link rel="stylesheet" href="{pre}assets/css/frame.css">
<link rel="stylesheet" href="{pre}assets/css/wow.css">
<link rel="stylesheet" href="{pre}assets/css/footer.css">
<link rel="stylesheet" href="{pre}assets/css/lab.css">
<link rel="stylesheet" href="{pre}guide/guide.css">
<link rel="stylesheet" href="{pre}assets/css/float.css">
</head>
<body class="lab-body v4" data-root="{pre}"{room_attr}>
<a class="skip" href="#main">Skip to content</a>
<header class="lab-top topbar">
<div class="bar-inner">
<button type="button" class="menu-btn" id="menu-btn" data-tour="menu" aria-controls="edge-menu" aria-expanded="false">Menu</button>
<a class="brand" href="{pre}index.html">{monkey_icon()}<span>CypherMonkey<small>AI omni-channel lab</small></span></a>
<p class="identity-line">A lifelong-learner lab on AI omni-channel agents, including voice and the contact center. Free to read. Nothing sold.</p>
<button type="button" class="display-btn" id="display-btn" aria-controls="display-panel" aria-expanded="false">Display</button>
</div>
<div id="display-panel" class="display-panel" hidden>
<p class="a11y-label">Font size</p>
<div class="a11y-font" role="group" aria-label="Font size">
<button type="button" class="themed" data-font-set="sm">A-</button>
<button type="button" class="themed" data-font-set="md">A</button>
<button type="button" class="themed" data-font-set="lg">A+</button>
<button type="button" class="themed" data-font-set="xl">A++</button>
</div>
<p class="a11y-label">Font colour</p>
<div class="ink-row" role="group" aria-label="Font colour">
<button type="button" class="ink-swatch" data-ink="#16182b" style="background:#16182b" aria-label="Near black"></button>
<button type="button" class="ink-swatch" data-ink="#0f172a" style="background:#0f172a" aria-label="Slate"></button>
<button type="button" class="ink-swatch" data-ink="#082f49" style="background:#082f49" aria-label="Deep blue"></button>
<button type="button" class="ink-swatch" data-ink="#7c2d12" style="background:#7c2d12" aria-label="Brown"></button>
<button type="button" class="ink-swatch" data-ink="#134e4a" style="background:#134e4a" aria-label="Teal"></button>
</div>
<label class="a11y-label">Custom font colour <input data-ink-input type="color" value="#16182b" aria-label="Custom font colour"></label>
<p class="a11y-label">Theme</p>
<div class="a11y-themes" role="radiogroup" aria-label="Theme">
<button type="button" class="themed" role="radio" data-set-theme="bright">Bright</button>
<button type="button" class="themed" role="radio" data-set-theme="dark">Dark</button>
<button type="button" class="themed" role="radio" data-set-theme="contrast">Contrast</button>
<button type="button" class="themed" role="radio" data-set-theme="holographic">Holographic</button>
<button type="button" class="themed" role="radio" data-set-theme="cyberpunk">Cyberpunk</button>
<button type="button" class="themed" role="radio" data-set-theme="neon">Neon</button>
<button type="button" class="themed" role="radio" data-set-theme="sunset">Sunset</button>
</div>
<button type="button" class="themed" data-motion-toggle aria-pressed="false">Reduce motion</button>
</div>
</header>
<div class="canon-banner" aria-label="Sections">
<button type="button" class="rail-arrow" data-rail="-1" aria-label="Scroll sections backward">‹</button>
<nav class="section-rail" aria-label="Sections">{''.join(chips)}</nav>
<button type="button" class="rail-arrow" data-rail="1" aria-label="Scroll sections forward">›</button>
</div>
<div id="edge-scrim" class="edge-scrim"></div>
<aside id="edge-menu" class="edge-menu" aria-label="Menu">
<div class="edge-menu-bar"><p>Menu</p><button type="button" class="edge-close" aria-label="Close menu">Close</button></div>
<label class="a11y-label" for="edge-search">Search the menu</label>
<input id="edge-search" class="edge-search" type="search" placeholder="Search" aria-label="Search the menu">
<ul class="edge-list">{''.join(drawer)}</ul>
<p class="edge-desc">Prior study notes stay in the campus shell. This front door is the learning lab.</p>
</aside>
<main id="main" class="wrap">{main}</main>
{footer()}
{dock}
<div id="guide-mount"></div>
<script src="{pre}assets/js/pages.js" defer></script>
<script src="{pre}assets/js/frame.js" defer></script>
<script src="{pre}assets/js/float.js" defer></script>
<script src="{pre}assets/js/university.js" defer></script>
<script src="{pre}guide/guide.js" defer></script>
</body>
</html>
"""


def buttons(links: list[tuple[str, str, str]]) -> str:
    parts = []
    for href, kind, label in links:
        parts.append(f'<a class="dl-btn dl-{esc(kind)}" href="{esc(href)}">{esc(label)} {esc(kind)}</a>')
    return f'<p class="export-row">{"".join(parts)}</p>'


def disc_svg_footer(y: int) -> str:
    line = "Independent, free, learning-only student project. Not affiliated or endorsed. Trademarks belong to their owners. Nothing sold."
    return (
        f'<text x="24" y="{y}" fill="#0f172a" font-size="14" font-family="system-ui,sans-serif">{esc(line)}</text>'
        f'<text x="24" y="{y + 22}" fill="#0c4a6e" font-size="13" font-family="system-ui,sans-serif">'
        f'{esc(SITE)}  ·  {esc(GITHUB)}  ·  {esc(LINKTREE)}</text>'
    )


def arch_svg(title: str, actors: list[str]) -> str:
    colors = ["#22d3ee", "#818cf8", "#f472b6", "#fbbf24"]
    boxes = []
    x0 = 36
    for index, name in enumerate(actors[:4]):
        x = x0 + index * 190
        fill = colors[index % 4]
        boxes.append(
            f'<rect x="{x}" y="120" width="160" height="72" rx="16" fill="{fill}" stroke="#0f172a" stroke-width="3"/>'
            f'<text x="{x + 80}" y="162" text-anchor="middle" fill="#0f172a" font-size="16" font-family="system-ui,sans-serif" font-weight="700">{esc(name)}</text>'
        )
        if index:
            boxes.append(f'<path d="M{x - 22} 156 H{x - 4}" stroke="#0f172a" stroke-width="3" marker-end="url(#arrow)"/>')
    body = "".join(boxes)
    return f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 820 280" role="img">
<title>{esc(title)}</title>
<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="6" refY="3" orient="auto"><path d="M0 0 L6 3 L0 6 Z" fill="#0f172a"/></marker></defs>
<rect width="820" height="280" fill="#fff7ed"/>
<text x="24" y="42" fill="#0f172a" font-size="22" font-family="system-ui,sans-serif" font-weight="700">{esc(title)}</text>
{body}
{disc_svg_footer(236)}
</svg>'''


def flow_svg(title: str, steps: list[str]) -> str:
    count = max(len(steps), 1)
    width = 860
    gap = (width - 80) / count
    parts = []
    points = []
    for index, step in enumerate(steps):
        cx = 50 + gap * index + gap / 2
        points.append((cx, 130))
        parts.append(
            f'<circle cx="{cx:.0f}" cy="130" r="36" fill="#fde68a" stroke="#b45309" stroke-width="3"/>'
            f'<text x="{cx:.0f}" y="126" text-anchor="middle" fill="#0f172a" font-size="13" font-family="system-ui,sans-serif" font-weight="700">{index + 1}</text>'
            f'<text x="{cx:.0f}" y="190" text-anchor="middle" fill="#0f172a" font-size="13" font-family="system-ui,sans-serif">{esc(step[:28])}</text>'
        )
    path = " ".join(f"{'M' if i == 0 else 'L'} {x:.0f} {y}" for i, (x, y) in enumerate(points))
    motion = f'<circle r="9" fill="#e11d48"><animateMotion dur="4s" repeatCount="indefinite" path="{path}"/></circle>' if len(points) > 1 else ""
    return f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} 280" role="img" data-animated="1">
<title>{esc(title)}</title>
<rect width="{width}" height="280" fill="#fff7ed"/>
<text x="24" y="40" fill="#0f172a" font-size="22" font-family="system-ui,sans-serif" font-weight="700">{esc(title)}</text>
<path d="{path}" fill="none" stroke="#0e7490" stroke-width="4"/>
{''.join(parts)}
{motion}
{disc_svg_footer(236)}
</svg>'''


def sheet_svg(title: str, lines: list[str]) -> str:
    rows = []
    y = 92
    for line in lines:
        size = 14 if line.startswith("http") or line.startswith("last checked") else 18
        rows.append(
            f'<text x="36" y="{y}" fill="#0f172a" font-size="{size}" font-family="system-ui,sans-serif">{esc(line[:140])}</text>'
        )
        y += 28 if size == 14 else 36
    height = max(420, y + 90)
    return f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 800 {height}" role="img">
<title>{esc(title)}</title>
<rect width="800" height="{height}" fill="#fff7ed"/>
<rect x="20" y="20" width="760" height="{height - 40}" rx="18" fill="#ffffff" stroke="#0e7490" stroke-width="3"/>
<text x="36" y="64" fill="#0f172a" font-size="26" font-family="system-ui,sans-serif" font-weight="700">{esc(title)}</text>
{''.join(rows)}
{disc_svg_footer(height - 56)}
</svg>'''


def standalone(title: str, body: str, depth: int) -> str:
    pre = prefix(depth)
    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(title)}</title>
<link rel="stylesheet" href="{pre}assets/css/footer.css">
<link rel="stylesheet" href="{pre}assets/css/lab.css">
</head><body class="lab-body">
<main class="wrap"><h1>{esc(title)}</h1>
<p class="open-row"><a class="dl-btn" href="{pre}index.html">Open app</a>
<a class="dl-btn dl-HTML" href="{REPO}">View repo</a></p>
{body}
</main>
{footer()}
</body></html>
"""


def room_id(track_id: str, number: int) -> str:
    return f"{track_id}-R{number}"


def all_rooms():
    for track in TRACKS:
        for room in track["rooms"]:
            yield track, room


def build_diagrams() -> tuple[list[dict], list[dict]]:
    diagrams = []
    flows = []
    for track, room in all_rooms():
        rid = room_id(track["id"], room["n"])
        title = f'{rid} {room["title"]}'
        svg = arch_svg(title, room["actors"])
        rel = f"downloads/diagrams/{rid}-arch.svg"
        svg_png(rel, svg)
        html_rel = f"downloads/diagrams/{rid}-arch.html"
        img = f"{rid}-arch.png"
        keys = room_keys(rid, room["cites"])
        deeper = go_deeper_html(keys, f"../../rooms/sources/{rid}/sources.json")
        write_text(html_rel, standalone(title, f'<img class="diagram" src="{img}" alt="{esc(title)}"><p><a href="{rid}-arch.svg">SVG</a></p>{deeper}', 2), "HTML")
        diagrams.append({"id": rid, "title": title, "kind": "arch", "track": track["title"], "stem": f"{rid}-arch", "keys": keys})
        if room["flow"]:
            flow_title = f'{rid} flow'
            flow = flow_svg(flow_title, room["flow"])
            frel = f"downloads/diagrams/{rid}-flow.svg"
            svg_png(frel, flow)
            write_text(
                f"downloads/diagrams/{rid}-flow.html",
                standalone(flow_title, f'<img class="diagram" src="{rid}-flow.png" alt="{esc(flow_title)}"><p><a href="{rid}-flow.svg">SVG</a></p>{deeper}', 2),
                "HTML",
            )
            flows.append({"id": rid, "title": flow_title, "steps": room["flow"], "keys": keys})
    extras = [
        ("overview-omni", "One policy, three channels", ["Person", "Policy", "Voice", "Chat"]),
        ("overview-rag", "Retrieve then read", ["Question", "Index", "Passage", "Quote"]),
        ("overview-agent", "Capped tool loop", ["Goal", "Model", "Tool", "Stop"]),
        ("overview-mcp", "Host, client, server", ["Person", "Host", "Client", "Server"]),
        ("overview-handoff", "Handoff packet", ["Assistant", "Packet", "Person", "Stop"]),
        ("overview-latency", "Voice stages", ["Hear", "Think", "Speak", "Clock"]),
        ("overview-glossary", "Words in this lab", ["Token", "Chunk", "Tool", "Channel"]),
        ("overview-certs", "Certificates stay outside", ["Outline", "Fee", "Confirm", "No award"]),
        ("overview-voice-stack", "Speech, a model, speech", ["Audio", "STT", "Model", "Speech"]),
        ("overview-analytics", "Recording, then a label", ["Recording", "Transcript", "Label", "Rubric"]),
        ("overview-safety", "Packet, then a person", ["Assistant", "Packet", "Person", "Stop"]),
        ("overview-cc", "Person, voice, record, CRM", ["Person", "Voice", "Record", "CRM"]),
        ("overview-integrate", "Call, session, note, person", ["Call", "Session", "Note", "Person"]),
    ]
    for key, title, actors in extras:
        keys = OVERVIEW_KEYS[key]
        svg_png(f"downloads/diagrams/{key}.svg", arch_svg(title, actors))
        write_text(
            f"downloads/diagrams/{key}.html",
            standalone(title, f'<img class="diagram" src="{key}.png" alt="{esc(title)}">{go_deeper_html(keys)}', 2),
            "HTML",
        )
        diagrams.append({"id": key, "title": title, "kind": "overview", "track": "Overview", "stem": key, "keys": keys})
    for track in TRACKS:
        title = f'{track["id"]} map'
        keys: list[str] = []
        for room in track["rooms"]:
            for key in room_keys(room_id(track["id"], room["n"]), room["cites"]):
                if key not in keys:
                    keys.append(key)
        svg_png(f"downloads/diagrams/{track['id']}-map.svg", arch_svg(title, [room["title"].split()[0] for room in track["rooms"][:4]]))
        write_text(
            f"downloads/diagrams/{track['id']}-map.html",
            standalone(title, f'<img class="diagram" src="{track["id"]}-map.png" alt="{esc(title)}">{go_deeper_html(keys)}', 2),
            "HTML",
        )
        diagrams.append({"id": track["id"], "title": title, "kind": "map", "track": track["title"], "stem": f"{track['id']}-map", "keys": keys})
    return diagrams, flows


def export_section(rel: str, title: str, paragraphs: list[str], keys: list[str], json_href: str) -> None:
    body = "".join(f"<p>{esc(p)}</p>" for p in paragraphs)
    write_text(rel, standalone(title, body + go_deeper_html(keys, json_href), 2), "HTML")


def room_page(track: dict, room: dict) -> None:
    rid = room_id(track["id"], room["n"])
    pre = "../"
    five_rel = f"downloads/sections/{rid}-five.html"
    ex_rel = f"downloads/sections/{rid}-example.html"
    quiz_rel = f"downloads/sections/{rid}-quiz.html"
    src_rel = f"downloads/sections/{rid}-sources.html"
    eco_rel = f"downloads/sections/{rid}-ecosystem.html"
    keys = room_keys(rid, room["cites"])
    json_rel = write_room_sources(rid, room["title"], keys)
    json_href = "../../" + json_rel
    export_section(five_rel, f"{rid} in five minutes", room["five"] + [room["idea"]], keys, json_href)
    export_section(ex_rel, room["ex_title"], room["steps"], keys, json_href)
    quiz_lines = []
    for question, choices, answer, why in room["quiz"]:
        quiz_lines.append(question + " Answer: " + choices[answer] + ". " + why)
    export_section(quiz_rel, f"{rid} quiz", quiz_lines, keys, json_href)
    export_section(src_rel, f"{rid} sources", ["The links below are the sources for this room."], keys, json_href)
    eco_lines = [f"{term}: {meaning}" for term, meaning in room["terms"]]
    eco_lines += track["sisters"] + track["bridges"] + [track["future"]] + room["config"]
    eco_lines += [f"{problem} {fix}" for problem, fix in room["trouble"]]
    export_section(eco_rel, f"{rid} ecosystem", eco_lines, keys, json_href)
    fields = []
    for index, (question, choices, answer, why) in enumerate(room["quiz"]):
        labels = "".join(
            f'<label><input type="radio" name="q{index}" value="{choice_index}" required> {esc(choice)}</label>'
            for choice_index, choice in enumerate(choices)
        )
        fields.append(f'<fieldset data-answer="{answer}"><legend>{esc(question)}</legend>{labels}<p>{esc(why)}</p></fieldset>')
    flow_btns = ""
    if room["flow"]:
        flow_btns = buttons([
            (f"{pre}downloads/diagrams/{rid}-flow.svg", "SVG", "Flow"),
            (f"{pre}downloads/diagrams/{rid}-flow.png", "PNG", "Flow"),
            (f"{pre}downloads/diagrams/{rid}-flow.html", "HTML", "Flow"),
        ])
    main = f"""
<p class="kicker">{esc(track['title'])} · {esc(room['level'])} · {room['mins']} min</p>
<h1>{esc(room['title'])}</h1>
<div class="room-layout">
<nav aria-label="On this page"><ul class="toc">
<li><a href="#five">Five minutes</a></li>
<li><a href="#visual">Visual</a></li>
<li><a href="#example">Worked example</a></li>
<li><a href="#quiz">Quiz</a></li>
<li><a href="#ecosystem">Ecosystem</a></li>
<li><a href="#sources">Sources and go deeper</a></li>
</ul></nav>
<div>
<section class="section-block" id="five" data-read data-menu="learn"><h2>Learn in five minutes</h2>
<ul>{''.join(f'<li>{esc(item)}</li>' for item in room['five'])}</ul>
<p>{esc(room['idea'])}</p>
{buttons([(five_rel.replace('downloads/', pre + 'downloads/'), 'HTML', 'This section')])}
</section>
<section class="section-block" id="visual" data-read data-menu="diagrams"><h2>Visual</h2>
<img class="diagram" src="{pre}downloads/diagrams/{rid}-arch.png" alt="{esc(rid)} diagram">
{buttons([
    (f'{pre}downloads/diagrams/{rid}-arch.svg', 'SVG', 'Diagram'),
    (f'{pre}downloads/diagrams/{rid}-arch.png', 'PNG', 'Diagram'),
    (f'{pre}downloads/diagrams/{rid}-arch.html', 'HTML', 'Diagram'),
])}
{flow_btns}
</section>
<section class="section-block" id="example" data-read data-menu="learn"><h2>{esc(room['ex_title'])}</h2>
<ol>{''.join(f'<li>{esc(step)}</li>' for step in room['steps'])}</ol>
{buttons([(ex_rel.replace('downloads/', pre + 'downloads/'), 'HTML', 'This section')])}
</section>
<section class="section-block" id="quiz" data-read data-menu="play"><h2>Quiz</h2>
<form class="quiz">{''.join(fields)}<button type="submit" class="themed">Check</button><p class="quiz-result" aria-live="polite"></p></form>
{buttons([(quiz_rel.replace('downloads/', pre + 'downloads/'), 'HTML', 'This section')])}
</section>
<section class="section-block" id="ecosystem" data-read data-menu="learn"><h2>Ecosystem</h2>
<dl>{''.join(f'<dt>{esc(term)}</dt><dd>{esc(meaning)}</dd>' for term, meaning in room['terms'])}</dl>
<p>Sisters: {esc(', '.join(track['sisters']))}.</p>
<p>{esc(track['bridges'][0])}</p>
<p>{esc(track['future'])}</p>
<h3>Config</h3><ul>{''.join(f'<li>{esc(item)}</li>' for item in room['config'])}</ul>
<h3>Troubleshooting</h3><ul>{''.join(f'<li>{esc(problem)} {esc(fix)}</li>' for problem, fix in room['trouble'])}</ul>
{buttons([(eco_rel.replace('downloads/', pre + 'downloads/'), 'HTML', 'This section')])}
</section>
{buttons([(src_rel.replace('downloads/', pre + 'downloads/'), 'HTML', 'Sources')])}
{go_deeper_html(keys, pre + json_rel, "sources")}
</div></div>
"""
    current = {"T9": "voice", "T10": "analytics", "T11": "safety", "T12": "ecosystem"}.get(track["id"], "learn")
    write_text(f"rooms/{rid}.html", chrome(f"{room['title']} · AI omni-channel lab", main, 1, current, rid), "HTML")


HUB_HREF = {
    "T9": "voice-stack.html",
    "T10": "speech-analytics.html",
    "T11": "agent-safety.html",
    "T12": "ecosystem.html",
}


def track_keys(track: dict) -> list[str]:
    keys: list[str] = []
    for room in track["rooms"]:
        for key in room_keys(room_id(track["id"], room["n"]), room["cites"]):
            if key not in keys:
                keys.append(key)
    return keys


def match_board(pairs: list[tuple[str, str]]) -> str:
    left = "".join(
        f'<button type="button" class="themed" data-pair="{index}" data-side="a">{esc(left_label)}</button>'
        for index, (left_label, _right) in enumerate(pairs)
    )
    right = "".join(
        f'<button type="button" class="themed" data-pair="{index}" data-side="b">{esc(right_label)}</button>'
        for index, (_left, right_label) in reversed(list(enumerate(pairs)))
    )
    return f'<div data-match><div>{left}</div><div>{right}</div><p data-match-status aria-live="polite"></p></div>'


def sort_board(steps: list[str]) -> str:
    items = "".join(
        f'<li data-step="{index}"><span>{esc(step)}</span> '
        f'<button type="button" class="themed" data-move="up">Up</button> '
        f'<button type="button" class="themed" data-move="down">Down</button></li>'
        for index, step in enumerate(reversed(steps))
    )
    order = ",".join(str(index) for index in range(len(steps)))
    return (
        f'<ul data-sort data-order="{order}">{items}</ul>'
        f'<button type="button" class="themed" data-sort-check>Check order</button>'
        f'<p data-sort-status aria-live="polite"></p>'
    )


def build_pages(diagrams: list[dict], flows: list[dict]) -> None:
    track_cards = []
    short = []
    for track in TRACKS:
        first = track["rooms"][0]
        rid = room_id(track["id"], first["n"])
        href = HUB_HREF.get(track["id"], f"rooms/{rid}.html")
        track_cards.append(
            f'<a class="track-card" href="{href}"><strong>{esc(track["title"])}</strong><span>{esc(track["blurb"])}</span></a>'
        )
        short.append(f'<li><a href="rooms/{rid}.html">{esc(first["title"])}</a> — {esc(first["five"][0])}</li>')

    def cite(key: str) -> str:
        item = entries([key])[0]
        if item["status"] == "opened" and item.get("url"):
            return f'<a href="{esc(item["url"])}">{esc(item["title"])}</a>'
        return f'{esc(item["title"])} (confirm)'

    home = f"""
<p class="kicker">Lifelong-learner lab</p>
<div class="hero"><div>
<h1>Learn <span class="hue">AI omni-channel agents</span> without a job claim.</h1>
<p class="lede">Programming, models, retrieval, agents, MCP, voice and messaging, a voice stack, speech analytics, agent safety, a contact-center map, and a cloud or cluster story you re-check. This is a free learning lab. It does not claim professional GenAI work.</p>
<div class="choice-row">
<a class="choice" href="learn.html#ten"><strong>Finish in 10 minutes</strong><span>The short path only.</span></a>
<a class="choice" href="learn.html"><strong>Full course</strong><span>{sum(len(t['rooms']) for t in TRACKS)} rooms across {len(TRACKS)} tracks.</span></a>
<a class="choice" href="downloads.html"><strong>Just download</strong><span>Diagrams, sheets, cards, and a zip.</span></a>
</div>
</div>
<svg class="illu hero-diagram" viewBox="0 0 680 320" role="img" aria-label="Four boxes: person, policy, voice, and chat"><defs><linearGradient id="hero-bg" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#082f49"/><stop offset="0.55" stop-color="#312e81"/><stop offset="1" stop-color="#4c0519"/></linearGradient><linearGradient id="hero-line" x1="0" y1="0" x2="1" y2="0"><stop offset="0" stop-color="#22d3ee"/><stop offset="0.5" stop-color="#c4b5fd"/><stop offset="1" stop-color="#fb7185"/></linearGradient><filter id="hero-glow" x="-40%" y="-40%" width="180%" height="180%"><feGaussianBlur stdDeviation="6" result="b"/><feMerge><feMergeNode in="b"/><feMergeNode in="SourceGraphic"/></feMerge></filter></defs><rect width="680" height="320" rx="28" fill="url(#hero-bg)"/><g opacity=".28" stroke="#67e8f9" stroke-width="1"><path d="M0 40H680M0 100H680M0 160H680M0 220H680M0 280H680M40 0V320M120 0V320M200 0V320M280 0V320M360 0V320M440 0V320M520 0V320M600 0V320"/></g><ellipse class="orbit" cx="340" cy="178" rx="250" ry="86" fill="none" stroke="#67e8f9" stroke-width="3" opacity=".8"/><ellipse class="orbit" cx="340" cy="178" rx="190" ry="58" fill="none" stroke="#f472b6" stroke-width="3" opacity=".75"/><text x="28" y="40" fill="#fde68a" font-size="20" font-family="system-ui,sans-serif" font-weight="800">One policy, three channels</text><path class="flow" d="M160 150H196 M316 150H352 M472 150H508" fill="none" stroke="url(#hero-line)" stroke-width="6" stroke-linecap="round" filter="url(#hero-glow)"/><g filter="url(#hero-glow)"><rect x="40" y="112" width="120" height="76" rx="22" fill="#22d3ee"/><text x="100" y="158" text-anchor="middle" fill="#082f49" font-size="20" font-family="system-ui,sans-serif" font-weight="800">Person</text><rect x="196" y="112" width="120" height="76" rx="22" fill="#a78bfa"/><text x="256" y="158" text-anchor="middle" fill="#1e1b4b" font-size="20" font-family="system-ui,sans-serif" font-weight="800">Policy</text><rect x="352" y="112" width="120" height="76" rx="22" fill="#fb7185"/><text x="412" y="158" text-anchor="middle" fill="#4c0519" font-size="20" font-family="system-ui,sans-serif" font-weight="800">Voice</text><rect x="508" y="112" width="120" height="76" rx="22" fill="#fde047"/><text x="568" y="158" text-anchor="middle" fill="#713f12" font-size="20" font-family="system-ui,sans-serif" font-weight="800">Chat</text></g><circle class="pulse" cx="340" cy="248" r="16" fill="#22d3ee"/><circle class="pulse" cx="420" cy="248" r="12" fill="#f9a8d4"/><circle class="pulse" cx="250" cy="248" r="12" fill="#fde047"/></svg>
</div>
<h2>Tracks</h2>
<div class="track-grid">{''.join(track_cards)}</div>
<section class="section-block" id="what" data-read data-menu="home"><h2>What this lab means by an agent</h2>
<p>An omni-channel agent, here, is one policy that can meet a person on voice, on chat, or in a message, and can hand that person to a human. This site is a free learning lab. It does not claim a deployment, a job, or a certificate.</p>
<p>The {cite("nist")} is the vocabulary this lab uses for that risk. It is not a product certificate. The skip link, the menu, and the type-size control follow the intent of {cite("wcag")}. {cite("wiki-a11y")} is the short public overview of that idea.</p>
</section>
<section class="section-block" id="model" data-read data-menu="home"><h2>A model continues text</h2>
<p>{cite("attn")} describes the transformer architecture later language models build on. {cite("gpt3")} reports that a large model can follow a few examples written in the prompt. A fluent sentence can still be false. The check in this lab is the source, not the tone. {cite("wiki-llm")} and {cite("wiki-attn")} are the public summaries.</p>
</section>
<section class="section-block" id="retrieve" data-read data-menu="home"><h2>Retrieve, then answer</h2>
<p>{cite("rag")} puts retrieved passages into the generation step. The rule on this site is shorter: if the note is silent, refuse. Fluency is not a citation. {cite("wiki-rag")} is the public summary of that pattern.</p>
</section>
<section class="section-block" id="loop" data-read data-menu="home"><h2>A loop is not a workflow</h2>
<p>{cite("react")} interleaves a reason step and a tool step. {cite("agents")} is the vendor essay this lab uses when it separates a fixed workflow from a loop that chooses a tool. Keep a schema, a step cap, and a trace. A person approves an action that is hard to undo.</p>
</section>
<section class="section-block" id="protocol" data-read data-menu="home"><h2>One way to name tools</h2>
<p>{cite("mcp")} names a host, a client, and a server. {cite("mcp-docs")} is the public introduction. {cite("wiki-mcp")} is the encyclopedia summary. Confirm the current revision before you build. This page opens no port.</p>
<p>The older campus shell, with 593 study pages, stays available as prior notes at <a href="campus/ai-lab-university.html">campus/ai-lab-university.html</a>.</p>
</section>
{go_deeper_html(["nist", "wcag", "wiki-a11y", "attn", "gpt3", "wiki-attn", "wiki-llm", "rag", "wiki-rag", "react", "agents", "mcp", "mcp-docs", "wiki-mcp"], "", "sources")}
"""
    write_text("index.html", chrome("AI omni-channel learning lab", home, 0, "home"), "HTML")
    learn_blocks = []
    for track in TRACKS:
        items = "".join(
            f'<li><a href="rooms/{room_id(track["id"], room["n"])}.html">{esc(room["title"])}</a> · {esc(room["level"])} · {room["mins"]} min</li>'
            for room in track["rooms"]
        )
        chapter_keys: list[str] = []
        json_links = []
        for room in track["rooms"]:
            rid = room_id(track["id"], room["n"])
            for key in room_keys(rid, room["cites"]):
                if key not in chapter_keys:
                    chapter_keys.append(key)
            json_links.append(f'<a href="rooms/sources/{rid}/sources.json">{esc(rid)} sources.json</a>')
        hub = HUB_HREF.get(track["id"])
        heading = f'<a href="{hub}">{esc(track["title"])}</a>' if hub else esc(track["title"])
        learn_blocks.append(
            f'<section class="section-block chapter" id="chapter-{esc(track["id"])}" data-read data-menu="learn">'
            f'<h2>{heading}</h2><p>{esc(track["blurb"])}</p><ol>{items}</ol>'
            f'{go_deeper_html(chapter_keys, "", "sources-" + track["id"])}'
            f'<p>Room files: {" · ".join(json_links)}</p></section>'
        )
    opening_keys: list[str] = []
    for track in TRACKS:
        first = track["rooms"][0]
        for key in room_keys(room_id(track["id"], first["n"]), first["cites"]):
            if key not in opening_keys:
                opening_keys.append(key)
    learn = f"""
<h1>Learn</h1>
<section class="section-block chapter" id="ten" data-read data-menu="learn"><h2>Finish in 10 minutes</h2>
<p>Read only these opening rooms, export one diagram, and stop.</p>
<ol>{''.join(short)}</ol>
{buttons([('downloads/sheets/five-minute-path.html','HTML','Short path'),('downloads/sheets/five-minute-path.pdf','PDF','Short path'),('downloads/sheets/five-minute-path.png','PNG','Short path')])}
{go_deeper_html(opening_keys, "", "sources-ten")}
</section>
{''.join(learn_blocks)}
"""
    write_text("learn.html", chrome("Learn · AI omni-channel lab", learn, 0, "learn"), "HTML")
    cards = []
    for item in diagrams:
        cards.append(
            f'<article class="card"><h2>{esc(item["title"])}</h2><p>{esc(item["track"])}</p>'
            f'<img class="diagram" src="downloads/diagrams/{item["stem"]}.png" alt="{esc(item["title"])}">'
            f'{buttons([(f"downloads/diagrams/{item["stem"]}.svg", "SVG", "File"), (f"downloads/diagrams/{item["stem"]}.png", "PNG", "File"), (f"downloads/diagrams/{item["stem"]}.html", "HTML", "File")])}'
            f'{go_deeper_html(item["keys"])}'
            f'</article>'
        )
    write_text("diagrams.html", chrome("Diagrams · AI omni-channel lab", f'<h1>Diagrams</h1><p data-read data-menu="diagrams">{len(diagrams)} drawings. Each one has SVG, PNG, and HTML, and the disclaimer is in the picture.</p><div class="card-grid">{"".join(cards)}</div>', 0, "diagrams"), "HTML")
    flow_cards = []
    for item in flows:
        flow_cards.append(
            f'<article class="card"><h2>{esc(item["title"])}</h2>'
            f'<img class="diagram" src="downloads/diagrams/{item["id"]}-flow.png" alt="{esc(item["title"])}">'
            f'{buttons([(f"downloads/diagrams/{item["id"]}-flow.svg","SVG","Flow"),(f"downloads/diagrams/{item["id"]}-flow.png","PNG","Flow"),(f"downloads/diagrams/{item["id"]}-flow.html","HTML","Flow")])}'
            f'{go_deeper_html(item["keys"])}'
            f'</article>'
        )
    write_text("flows.html", chrome("Flows · AI omni-channel lab", f'<h1>Flows</h1><p data-read>{len(flows)} animated flows. The PNG is one frame. The SVG keeps the motion.</p><div class="card-grid">{"".join(flow_cards)}</div>', 0, "flows"), "HTML")


def build_chapter_pages() -> None:
    by_id = {track["id"]: track for track in TRACKS}
    voice = by_id["T9"]
    analytics = by_id["T10"]
    safety = by_id["T11"]
    voice_rooms = "".join(
        f'<li><a href="rooms/{room_id("T9", room["n"])}.html">{esc(room["title"])}</a></li>'
        for room in voice["rooms"]
    )
    analytics_rooms = "".join(
        f'<li><a href="rooms/{room_id("T10", room["n"])}.html">{esc(room["title"])}</a></li>'
        for room in analytics["rooms"]
    )
    safety_rooms = "".join(
        f'<li><a href="rooms/{room_id("T11", room["n"])}.html">{esc(room["title"])}</a></li>'
        for room in safety["rooms"]
    )
    voice_page = f"""
<p class="kicker">Voice AI agent stack</p>
<h1>Voice stack</h1>
<p class="lede">Speech in, a model, speech out, and the network that carries the audio. This lab places no call and sends no audio.</p>
<section class="section-block" id="picture" data-read data-menu="voice"><h2>The cascaded row</h2>
<img class="diagram" src="downloads/diagrams/overview-voice-stack.png" alt="Four boxes: audio, speech to text, a model, and speech">
{buttons([
    ("downloads/diagrams/overview-voice-stack.svg", "SVG", "Diagram"),
    ("downloads/diagrams/overview-voice-stack.png", "PNG", "Diagram"),
    ("downloads/diagrams/overview-voice-stack.html", "HTML", "Diagram"),
    ("downloads/diagrams/T9-map.html", "HTML", "Track map"),
])}
<p>LiveKit's pipeline page names a cascaded pipeline as speech-to-text, then a language model, then text-to-speech. A realtime model, on that page, consumes and produces speech in one model. Pipecat's overview lists a transport, speech recognition, a language model, speech synthesis, and a transport again.</p>
</section>
<section class="section-block" id="sort" data-read data-menu="play"><h2>Sort the cascaded order</h2>
<p>Put the first voice room's flow back in order. A speech-to-speech model is a different shape. It does not sit inside this row.</p>
{sort_board(["Hear audio", "Transcribe", "Decide", "Speak"])}
</section>
<section class="section-block" id="match" data-read data-menu="play"><h2>Match the words</h2>
{match_board([
    ("STT", "Speech to text"),
    ("TTS", "Text to speech"),
    ("VAD", "Detects speech versus silence"),
    ("Barge-in", "The person speaks while the reply still plays"),
])}
</section>
<section class="section-block" id="rooms" data-read data-menu="voice"><h2>Rooms</h2>
<ol>{voice_rooms}</ol>
<p>Latency sentences stay attached to the page that wrote them. This lab states no service-level agreement and no price.</p>
</section>
{go_deeper_html(track_keys(voice), "", "sources")}
"""
    analytics_page = f"""
<p class="kicker">Speech analytics and quality</p>
<h1>Analytics</h1>
<p class="lede">A recording, a transcript, a label, and a rubric. A fluent sentence is not the check.</p>
<section class="section-block" id="picture" data-read data-menu="analytics"><h2>After the call</h2>
<img class="diagram" src="downloads/diagrams/overview-analytics.png" alt="Four boxes: recording, transcript, label, and rubric">
{buttons([
    ("downloads/diagrams/overview-analytics.svg", "SVG", "Diagram"),
    ("downloads/diagrams/overview-analytics.png", "PNG", "Diagram"),
    ("downloads/diagrams/overview-analytics.html", "HTML", "Diagram"),
    ("downloads/diagrams/T10-map.html", "HTML", "Track map"),
])}
<p>Wikipedia defines speech analytics as analyzing recorded calls. The live reply and that later reading are different clocks.</p>
</section>
<section class="section-block" id="match" data-read data-menu="play"><h2>Match the measure</h2>
{match_board([
    ("Precision", "Share of search results that are relevant"),
    ("Recall", "Share of relevant items the search returned"),
    ("Sentiment", "A polarity label, not a stated fact"),
    ("Redaction", "Mask a name and a linked identifier"),
])}
</section>
<section class="section-block" id="rooms" data-read data-menu="analytics"><h2>Rooms</h2>
<ol>{analytics_rooms}</ol>
<p>Do not paste a real customer's recording into this lab.</p>
</section>
{go_deeper_html(track_keys(analytics), "", "sources")}
"""
    safety_page = f"""
<p class="kicker">AI agent safety</p>
<h1>Safety</h1>
<p class="lede">A person can take the turn. A guardrail is a check you can fail. Recording notes on this site are educational and are not legal advice.</p>
<section class="section-block" id="picture" data-read data-menu="safety"><h2>The packet</h2>
<img class="diagram" src="downloads/diagrams/overview-safety.png" alt="Four boxes: assistant, packet, person, and stop">
{buttons([
    ("downloads/diagrams/overview-safety.svg", "SVG", "Diagram"),
    ("downloads/diagrams/overview-safety.png", "PNG", "Diagram"),
    ("downloads/diagrams/overview-safety.html", "HTML", "Diagram"),
    ("downloads/diagrams/T11-map.html", "HTML", "Track map"),
])}
<p>A human handoff in this lab means a person continues the turn. Pipecat's overview uses the word handoff for control moving to another agent. That is a different sentence.</p>
</section>
<section class="section-block" id="sort" data-read data-menu="play"><h2>Sort the packet</h2>
<p>The handoff room names these four lines. Put them back in that order.</p>
{sort_board(["Write the channel", "Write what was said", "Write the note id", "Write the open question"])}
</section>
<section class="section-block" id="rooms" data-read data-menu="safety"><h2>Rooms</h2>
<ol>{safety_rooms}</ol>
<p>This chapter states no legal conclusion and configures no emergency calling.</p>
</section>
{go_deeper_html(track_keys(safety), "", "sources")}
"""
    write_text("voice-stack.html", chrome("Voice stack · AI omni-channel lab", voice_page, 0, "voice"), "HTML")
    write_text("speech-analytics.html", chrome("Analytics · AI omni-channel lab", analytics_page, 0, "analytics"), "HTML")
    write_text("agent-safety.html", chrome("Safety · AI omni-channel lab", safety_page, 0, "safety"), "HTML")


def build_play_listen_help() -> None:
    track, room = TRACKS[0], TRACKS[0]["rooms"][0]
    sort_steps = room["flow"] or ["Read", "Export", "Stop"]
    items = "".join(
        f'<li data-step="{index}"><span>{esc(step)}</span> <button type="button" class="themed" data-move="up">Up</button> <button type="button" class="themed" data-move="down">Down</button></li>'
        for index, step in enumerate(reversed(sort_steps))
    )
    order = ",".join(str(index) for index in range(len(sort_steps)))
    pairs = [("Token", "A piece a model reads"), ("Tool", "A function with a schema"), ("Channel", "Voice, chat, or message")]
    left = "".join(f'<button type="button" data-pair="{i}" data-side="a">{esc(a)}</button>' for i, (a, _b) in enumerate(pairs))
    right = "".join(f'<button type="button" data-pair="{i}" data-side="b">{esc(b)}</button>' for i, (_a, b) in reversed(list(enumerate(pairs))))
    play = f"""
<h1>Play</h1>
<section class="section-block" data-read data-menu="play"><h2>Match</h2>
<div data-match><div>{left}</div><div>{right}</div><p data-match-status aria-live="polite"></p></div>
</section>
<section class="section-block"><h2>Sort the first flow</h2>
<ul data-sort data-order="{order}">{items}</ul>
<button type="button" class="themed" data-sort-check>Check order</button>
<p data-sort-status aria-live="polite"></p>
</section>
<section class="section-block"><h2>Scenario</h2>
<form class="scenario" data-answer="1">
<p>A note has no price. The model offers a number. What do you keep?</p>
<label><input type="radio" name="s" value="0" required> The number, because it is fluent.</label>
<label><input type="radio" name="s" value="1"> The refusal, because the note is silent.</label>
<button type="submit" class="themed">Check</button>
<p class="quiz-result" aria-live="polite"></p>
</form>
<p>More checks live in each room. This page is practice, not a certificate.</p>
</section>
{go_deeper_html(["nist", "wiki-llm"], "", "sources")}
"""
    write_text("play.html", chrome("Play · AI omni-channel lab", play, 0, "play"), "HTML")
    listen = f"""
<h1>Listen</h1>
<section class="section-block" data-read data-menu="listen"><h2>Read-along</h2>
<p>Open the guide and choose Read this page. The spoken sentence is the block on screen. The matching menu chip glows.</p>
<p>The browser picks the most natural English voice it has. You can change voice, rate, and pitch. Ambient sound stays off until you turn it on.</p>
<p>If the microphone is missing, type instead. This lab ships no pre-recorded speech clips.</p>
</section>
<section class="section-block" data-read data-menu="learn"><h2>A short passage</h2>
<p>An agent here is a loop that may call a tool. A workflow is enough when the steps never change. A voice turn still uses that cap.</p>
<p>Retrieval puts a passage in the call. If the passage is missing, refuse. Fluency is not a citation.</p>
</section>
{go_deeper_html(["mdn", "wiki-asr", "wiki-tts"], "", "sources")}
"""
    write_text("listen.html", chrome("Listen · AI omni-channel lab", listen, 0, "listen"), "HTML")
    help_page = f"""
<h1>Help</h1>
<section class="section-block" data-read data-menu="help"><h2>Menu</h2>
<p>The Menu button is at the top left. The drawer slides in. Escape or a tap on the dim page closes it. On a phone, Menu is also the first item in the bottom bar.</p>
</section>
<section class="section-block" data-read><h2>Course</h2>
<p>Learn holds the ten-minute path and the full track list. Voice stack, Analytics, and Safety are their own menu items. Ecosystem holds the contact-center map. Each room has a five-minute block, a picture, an example, a quiz, sources, and an ecosystem panel.</p>
</section>
<section class="section-block" data-read><h2>Guide</h2>
<p>The monkey greets you in the centre once, then docks. On a phone the guide is a sheet above the bottom bar. Offers come first. Voice and read-along stay folded until you open them.</p>
<p><button type="button" class="themed" data-replay-tour>Replay the tour</button></p>
</section>
<section class="section-block" data-read data-menu="downloads"><h2>Downloads</h2>
<p>Every room has export buttons. The Downloads page lists the files and the headless test table. Open app returns here. View repo opens the GitHub repository.</p>
</section>
{go_deeper_html(["wcag", "wiki-a11y"], "", "sources")}
"""
    write_text("help.html", chrome("Help · AI omni-channel lab", help_page, 0, "help"), "HTML")
    legal = f"""
<h1>Legal and credits</h1>
<section class="section-block" data-read><h2>What this is</h2>
<p>{esc(DISC)}</p>
<p>Code in this repository is MIT. Lesson text is CC BY-NC-SA 4.0. No warranty.</p>
<p>Trademarks named in citations belong to their owners. Naming a paper or a product page is not an endorsement.</p>
<p>Contact for a correction or removal: <a href="{CONTACT}">the contact page</a>.</p>
</section>
{go_deeper_html(["cc", "mit", "usc-2511", "fcc-mlts"], "", "sources")}
"""
    write_text("legal.html", chrome("Legal · AI omni-channel lab", legal, 0, "legal"), "HTML")
    eco = []
    for track in TRACKS:
        eco.append(
            f'<section class="section-block" data-read><h2>{esc(track["title"])}</h2>'
            f'<p>Sisters: {esc(", ".join(track["sisters"]))}.</p>'
            f'<p>{esc(" ".join(track["bridges"]))}</p><p>{esc(track["future"])}</p></section>'
        )
    cc = next(track for track in TRACKS if track["id"] == "T12")
    cc_rooms = "".join(
        f'<li><a href="rooms/{room_id("T12", room["n"])}.html">{esc(room["title"])}</a></li>'
        for room in cc["rooms"]
    )
    ecosystem = f"""
<h1>Ecosystem</h1>
<p>A learning map of the systems beside a call. No vendor deployment is claimed. Notes on recording, consent, and emergency calling are educational and are not legal advice.</p>
<section class="section-block" id="map" data-read data-menu="ecosystem"><h2>Mind map</h2>
<img class="diagram" src="downloads/diagrams/overview-cc.png" alt="Four boxes: person, voice, record, and CRM">
{buttons([
    ("downloads/diagrams/overview-cc.svg", "SVG", "Mind map"),
    ("downloads/diagrams/overview-cc.png", "PNG", "Mind map"),
    ("downloads/diagrams/overview-cc.html", "HTML", "Mind map"),
    ("downloads/diagrams/T12-map.html", "HTML", "Track map"),
])}
<p>The person stays in the center. Voice, the recording, and the customer record sit beside the call. Emergency calling stays off the model arrow.</p>
</section>
<section class="section-block" id="integrate" data-read data-menu="ecosystem"><h2>Integration map</h2>
<img class="diagram" src="downloads/diagrams/overview-integrate.png" alt="Four boxes: call, session, note, and person">
{buttons([
    ("downloads/diagrams/overview-integrate.svg", "SVG", "Integration map"),
    ("downloads/diagrams/overview-integrate.png", "PNG", "Integration map"),
    ("downloads/diagrams/overview-integrate.html", "HTML", "Integration map"),
])}
<p>One arrow is the live reply. Another runs from a recording to a transcript and a rubric. A third carries the handoff packet to a person. Drawing an arrow does not configure a vendor.</p>
</section>
<section class="section-block" id="match" data-read data-menu="play"><h2>Match a sister to its job</h2>
{match_board([
    ("WFM", "Forecast, schedule, and attendance"),
    ("CRM", "Customer interactions, including the phone"),
    ("CPaaS", "Voice, video, and messaging in an app"),
    ("E911", "The FCC page on 911 for multi-line systems"),
])}
</section>
<section class="section-block" id="rooms" data-read data-menu="ecosystem"><h2>Contact-center rooms</h2>
<ol>{cc_rooms}</ol>
<p><a href="glossary.html">Glossary</a> · <a href="downloads/diagrams/overview-omni.html">One policy, three channels</a></p>
</section>
{''.join(eco)}
{go_deeper_html(track_keys(cc), "", "sources")}
"""
    write_text("ecosystem.html", chrome("Ecosystem · AI omni-channel lab", ecosystem, 0, "ecosystem"), "HTML")
    terms = []
    glossary_keys: list[str] = []
    for track, room in all_rooms():
        for key in room_keys(room_id(track["id"], room["n"]), room["cites"]):
            if key not in glossary_keys:
                glossary_keys.append(key)
        for term, meaning in room["terms"]:
            terms.append(f"<dt>{esc(term)}</dt><dd>{esc(meaning)}</dd>")
    write_text(
        "glossary.html",
        chrome(
            "Glossary · AI omni-channel lab",
            f"<h1>Glossary</h1><dl>{''.join(terms)}</dl>{go_deeper_html(glossary_keys, '', 'sources')}",
            0,
            "glossary",
        ),
        "HTML",
    )
    certs = f"""
<h1>Certifications</h1>
<section class="section-block" data-read><h2>Confirm before you pay</h2>
<p>Vendors and short-course sites publish certificates. The outline and the fee change. Confirm both on the issuer's own page before you pay.</p>
<p>This site awards none of those certificates. It does not claim that the author holds them. A folder with notes, a trace, a score, and a refusal is the portfolio this lab describes.</p>
</section>
{go_deeper_html(["nist", "cc"], "", "sources")}
"""
    write_text("certs.html", chrome("Certifications · AI omni-channel lab", certs, 0, "certs"), "HTML")
    labs = f"""
<h1>Python labs</h1>
<section class="section-block" data-read><h2>Zero-key, in this repository</h2>
<ul>
<li><code>python3 labs/01-local-hello.py</code> — local runner or an honest mock.</li>
<li><code>python3 labs/02-tool-call.py</code> — a sum from a tool.</li>
<li><code>python3 labs/03-golden-eval.py</code> — a fixed ten-question file.</li>
<li><code>python3 labs/06-rag-ablation.py</code> — one fixture, keyword versus a naive neighbor.</li>
<li><code>python3 labs/07-agent-loop.py</code> — a correction and a stop.</li>
<li><code>python3 labs/08-voice-latency-budget.py</code> — a stage p95 on toy milliseconds.</li>
</ul>
<p>No key is required. A missing local model is reported as a mock. These scripts are the worked examples the rooms point at. The milliseconds in lab 08 are a teaching fixture.</p>
</section>
{go_deeper_html(["py", "livekit-pipe", "pipecat", "xai-voice"], "", "sources")}
"""
    write_text("labs.html", chrome("Labs · AI omni-channel lab", labs, 0, "labs"), "HTML")


SHEETS = [
    ("five-minute-path", "Five-minute path", ["Open Learn and read only the opening room of each track.", "Export one SVG.", "Answer that room's three checks.", "Stop. The full course can wait.", "Progress stays in this browser."]),
    ("tokens-sampling", "Tokens and sampling", ["A token is not always a word.", "A space split is a toy. Label it.", "Temperature is not a truth dial.", "Record temperature, max tokens, and the stop.", "Confirm a model card before you quote a window."]),
    ("rag-checklist", "RAG checklist", ["Write the question.", "Retrieve passages.", "Read them before you answer.", "Quote or refuse.", "If notes disagree, say so.", "Date a stale note."]),
    ("agent-loop", "Agent loop", ["Name the goal.", "Validate tool arguments.", "Set a step cap.", "Stop when a result repeats.", "A person approves an undo-hard action.", "Keep the trace."]),
    ("mcp-map", "MCP map", ["Host: the app a person uses.", "Client: the protocol speaker.", "Server: tools, resources, or prompts.", "Consent names the action.", "Confirm the spec revision.", "This lab opens no port."]),
    ("voice-budget", "Voice budget", ["Hear, decide, speak.", "Speech is not the language model.", "Label millisecond figures as a fixture.", "This page states no SLA.", "The slow stage is the one to study.", "The tool cap still applies."]),
    ("handoff-packet", "Handoff packet", ["Channel.", "What was already said.", "Note id.", "The open question.", "No promise the note does not allow.", "If no person is available, say so."]),
    ("glossary-sheet", "Glossary card", ["Token: what the model reads.", "Chunk: what you index.", "Tool: a function with a schema.", "Channel: voice, chat, or message.", "Refusal: a successful silence.", "Confirm before relying: the honesty mark."]),
]

TEMPLATES = [
    ("lab-notebook", "Lab notebook", ["Date", "Command", "Input", "Output", "Mock or live", "What you will not claim"]),
    ("eval-card", "Eval card", ["Question", "Expected property", "Source note", "Settings", "Result", "One-line reason for a change"]),
    ("agent-trace", "Agent trace", ["Goal", "Proposal", "Schema check", "Tool result", "Stop reason", "Approver"]),
    ("channel-router", "Channel router", ["Policy", "Voice wording", "Chat wording", "Message wording", "Handoff packet", "Refusal"]),
    ("rag-checklist-form", "RAG checklist form", ["Corpus you may use", "Chunk note", "Retrieved passage", "Quote", "Refusal case", "Date of the note"]),
]


def build_downloads_assets() -> None:
    for slug, title, lines in SHEETS:
        keys = SHEET_KEYS[slug]
        drawn = lines + source_lines(keys)
        svg = sheet_svg(title, drawn)
        svg_png(f"downloads/sheets/{slug}.svg", svg)
        svg_pdf(f"downloads/sheets/{slug}.pdf", svg)
        body = "<ul>" + "".join(f"<li>{esc(line)}</li>" for line in lines) + "</ul>" + go_deeper_html(keys)
        write_text(f"downloads/sheets/{slug}.html", standalone(title, body, 2), "HTML")
    cards = []
    for track, room in all_rooms():
        rid = room_id(track["id"], room["n"])
        keys = room_keys(rid, room["cites"])
        for term, meaning in room["terms"]:
            cards.append({
                "front": term,
                "back": meaning,
                "room": rid,
                "sources": entries(keys),
            })
    payload = {"disclaimer": DISC, "contact": CONTACT, "last_checked": CHECKED, "cards": cards}
    write_text("downloads/cards/flashcards.json", json.dumps(payload, indent=2), "JSON")
    card_parts = []
    for card in cards:
        json_href = f'../../rooms/sources/{card["room"]}/sources.json'
        card_parts.append(
            f'<article class="card"><h2>{esc(card["front"])}</h2><p>{esc(card["back"])}</p><p>{esc(card["room"])}</p>'
            f'{go_deeper_html([item["id"] for item in card["sources"]], json_href)}'
            f'</article>'
        )
    card_html = "".join(card_parts)
    write_text("downloads/cards/flashcards.html", standalone("Flashcards", card_html, 2), "HTML")
    lines = [f'{card["front"]}: {card["back"]}' for card in cards]
    card_keys: list[str] = []
    for card in cards:
        for item in card["sources"]:
            if item["id"] not in card_keys:
                card_keys.append(item["id"])
    svg = sheet_svg("Flashcards", lines[:18] + source_lines(card_keys[:8]))
    svg_pdf("downloads/cards/flashcards.pdf", svg)
    for slug, title, fields in TEMPLATES:
        keys = TEMPLATE_KEYS[slug]
        svg = sheet_svg(title, [field + ": ____________________" for field in fields] + source_lines(keys))
        svg_pdf(f"downloads/templates/{slug}.pdf", svg)
        write_text(
            f"downloads/templates/{slug}.html",
            standalone(title, "<ul>" + "".join(f"<li>{esc(field)}</li>" for field in fields) + "</ul>" + go_deeper_html(keys), 2),
            "HTML",
        )
        document = Document()
        document.add_heading(title, 0)
        document.add_paragraph(DISC)
        document.add_paragraph(CONTACT)
        for field in fields:
            document.add_paragraph(field + ": ")
        document.add_heading("Sources and go deeper", level=1)
        for line in source_lines(keys)[1:]:
            document.add_paragraph(line)
        target = ROOT / "downloads" / "templates" / f"{slug}.docx"
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)
        RECORDS.append({"rel": f"downloads/templates/{slug}.docx", "kind": "DOCX", "bytes": target.stat().st_size, "disclaimer": True})
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1.2))
    box.text_frame.text = "AI omni-channel learning lab"
    note = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(3))
    note.text_frame.text = DISC
    for track in TRACKS:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1))
        title.text_frame.text = track["title"]
        run = title.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(32)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(9), Inches(4))
        chapter_keys: list[str] = []
        for room in track["rooms"]:
            for key in room_keys(room_id(track["id"], room["n"]), room["cites"]):
                if key not in chapter_keys:
                    chapter_keys.append(key)
        source_note = "\n".join(source_lines(chapter_keys[:6]))
        body.text_frame.text = track["blurb"] + "\n" + "\n".join(room["title"] for room in track["rooms"]) + "\n" + source_note
        foot = slide.shapes.add_textbox(Inches(0.5), Inches(6.4), Inches(9), Inches(0.8))
        foot.text_frame.text = "Learning-only student project. Nothing sold. " + LINKTREE
    pptx_path = ROOT / "downloads" / "ai-omni-lab-deck.pptx"
    deck.save(pptx_path)
    RECORDS.append({"rel": "downloads/ai-omni-lab-deck.pptx", "kind": "PPTX", "bytes": pptx_path.stat().st_size, "disclaimer": True})
    write_text("downloads/DISCLAIMER.txt", DISC + "\n" + CONTACT + "\n" + LINKTREE + "\n", "TXT")


def probe(record: dict) -> str:
    path = ROOT / record["rel"]
    data = path.read_bytes()[:8]
    kind = record["kind"]
    if kind == "PNG" and not data.startswith(b"\x89PNG"):
        return "PNG header missing"
    if kind == "PDF" and not data.startswith(b"%PDF"):
        return "PDF header missing"
    if kind in {"DOCX", "PPTX", "ZIP"}:
        try:
            with zipfile.ZipFile(path) as archive:
                archive.namelist()
        except zipfile.BadZipFile:
            return "zip open failed"
    if kind == "JSON":
        json.loads(path.read_text(encoding="utf-8"))
    if kind == "SVG" and b"<svg" not in path.read_bytes()[:200]:
        return "svg root missing"
    if kind == "HTML" and b"learning-only student project" not in path.read_bytes():
        return "disclaimer missing"
    if not record["disclaimer"] and kind in {"HTML", "SVG", "JSON", "TXT"}:
        return "disclaimer flag missing"
    return "opened headless"


def build_downloads_page() -> str:
    rows = []
    for record in RECORDS:
        if not record["rel"].startswith("downloads/"):
            continue
        status = probe(record)
        rows.append(
            f'<tr><td><a href="{esc(record["rel"].removeprefix("downloads/"))}">{esc(record["rel"])}</a></td>'
            f'<td>{esc(record["kind"])}</td><td>{record["bytes"]}</td>'
            f'<td>{"yes" if record["disclaimer"] else "no"}</td><td>{esc(status)}</td></tr>'
        )
    table = (
        '<div class="table-wrap"><table><thead><tr><th>File</th><th>Kind</th><th>Bytes</th><th>Disclaimer</th><th>Headless check</th></tr></thead><tbody>'
        + "".join(rows)
        + "</tbody></table></div>"
    )
    groups = [
        ("Diagrams", "downloads/diagrams/"),
        ("Sheets", "downloads/sheets/"),
        ("Flashcards", "downloads/cards/"),
        ("Templates", "downloads/templates/"),
    ]
    blocks = []
    for label, folder in groups:
        links = [f'<li><a href="{esc(record["rel"])}">{esc(Path(record["rel"]).name)}</a></li>' for record in RECORDS if record["rel"].startswith(folder) and record["kind"] in {"HTML", "PDF", "JSON", "DOCX"}]
        blocks.append(f'<section class="section-block" data-read data-menu="downloads"><h2>{label}</h2><p>{len(links)} files.</p><ul>{"".join(links)}</ul></section>')
    main = f"""
<h1>Downloads</h1>
<p class="open-row"><a class="dl-btn" href="index.html">Open app</a><a class="dl-btn dl-HTML" href="{REPO}">View repo</a>
<a class="dl-btn dl-ZIP" href="downloads/ai-omni-lab-offline.zip">Offline zip</a>
<a class="dl-btn dl-PPTX" href="downloads/ai-omni-lab-deck.pptx">Deck PPTX</a></p>
<p data-read data-menu="downloads">Every new drawing has a PNG. The zip is the new lab, not the older campus archive. Each file below repeats a sources block. Unopenable pages are marked confirm and are not linked.</p>
{go_deeper_html(["nist", "wcag", "attn", "rag", "react", "agents", "mcp", "mcp-docs", "mdn", "cc", "wiki-llm", "livekit-pipe", "xai-s2s", "wiki-analytics", "fcc-mlts"], "", "sources")}
{''.join(blocks)}
<h2>Headless test table</h2>
{table}
"""
    page = chrome("Downloads · AI omni-channel lab", main, 0, "downloads")
    write_text("downloads.html", page, "HTML")
    md = ["# Downloads", "", DISC, "", "| File | Kind | Bytes | Disclaimer | Check |", "|---|---|---:|---|---|"]
    for record in RECORDS:
        if record["rel"].startswith("downloads/"):
            md.append(f'| `{record["rel"]}` | {record["kind"]} | {record["bytes"]} | {"yes" if record["disclaimer"] else "no"} | {probe(record)} |')
    md.append("")
    md.append("The offline zip is `downloads/ai-omni-lab-offline.zip`. It is rebuilt after this table and excludes the older campus archive.")
    write_text("docs/DOWNLOADS.md", "\n".join(md) + "\n", "MD")
    return page


def build_shell_files(pages: list[dict]) -> None:
    payload = "window.AI_GUIDE_PAGES = " + json.dumps(pages) + ";\n"
    write_text("assets/js/pages.js", payload, "JS")
    favicon = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64">
<rect width="64" height="64" rx="14" fill="#fff7ed"/>
<circle cx="18" cy="24" r="8" fill="#f59e0b" stroke="#9a3412" stroke-width="2"/>
<circle cx="46" cy="24" r="8" fill="#f59e0b" stroke="#9a3412" stroke-width="2"/>
<circle cx="32" cy="36" r="16" fill="#fbbf24" stroke="#b45309" stroke-width="2"/>
<circle cx="26" cy="34" r="3" fill="#fff" stroke="#1c1408" stroke-width="1"/>
<circle cx="38" cy="34" r="3" fill="#fff" stroke="#1c1408" stroke-width="1"/>
<circle cx="27" cy="35" r="1.4" fill="#1c1408"/>
<circle cx="39" cy="35" r="1.4" fill="#1c1408"/>
<circle cx="26" cy="34" r="5" fill="none" stroke="#0e7490" stroke-width="1.6"/>
<circle cx="38" cy="34" r="5" fill="none" stroke="#0e7490" stroke-width="1.6"/>
<path d="M31 34 H33" stroke="#0e7490" stroke-width="1.4"/>
<path d="M27 42 Q32 46 37 42" fill="none" stroke="#1c1408" stroke-width="1.6" stroke-linecap="round"/>
</svg>'''
    write_text("favicon.svg", favicon, "SVG")
    for size, name in ((192, "icon-192.png"), (512, "icon-512.png")):
        png = cairosvg.svg2png(bytestring=favicon.encode("utf-8"), output_width=size, output_height=size)
        write_bytes(name, png, "PNG", False)
    manifest = {
        "name": "AI omni-channel learning lab",
        "short_name": "AI Lab",
        "start_url": "index.html",
        "display": "standalone",
        "background_color": "#fff7ed",
        "theme_color": "#fff7ed",
        "icons": [
            {"src": "icon-192.png", "sizes": "192x192", "type": "image/png"},
            {"src": "icon-512.png", "sizes": "512x512", "type": "image/png"},
        ],
    }
    write_text("manifest.webmanifest", json.dumps(manifest, indent=2), "JSON")
    core = [
        "index.html", "learn.html", "voice-stack.html", "speech-analytics.html", "agent-safety.html",
        "diagrams.html", "flows.html", "play.html", "listen.html",
        "downloads.html", "help.html", "legal.html", "ecosystem.html", "glossary.html", "certs.html", "labs.html",
        "favicon.svg", "manifest.webmanifest",
        "assets/css/frame.css", "assets/css/wow.css", "assets/css/footer.css", "assets/css/lab.css", "assets/css/float.css",
        "assets/js/frame.js", "assets/js/float.js", "assets/js/university.js", "assets/js/pages.js",
        "guide/guide.css", "guide/guide.js",
    ]
    core += [f"rooms/{room_id(track['id'], room['n'])}.html" for track, room in all_rooms()]
    files = json.dumps(core)
    sw = (
        "const CACHE='ai-omni-v4';const FILES=" + files + ";\n"
        "self.addEventListener('install',event=>event.waitUntil(caches.open(CACHE).then(cache=>cache.addAll(FILES))));\n"
        "self.addEventListener('activate',event=>event.waitUntil(self.clients.claim()));\n"
        "self.addEventListener('fetch',event=>{if(event.request.method!=='GET')return;"
        "event.respondWith(caches.match(event.request).then(hit=>hit||fetch(event.request)));});\n"
    )
    write_text("sw.js", sw, "JS")
    (ROOT / ".nojekyll").write_text("", encoding="utf-8")


def build_zip() -> None:
    destination = ROOT / "downloads" / "ai-omni-lab-offline.zip"
    include_roots = [
        "index.html", "learn.html", "voice-stack.html", "speech-analytics.html", "agent-safety.html",
        "diagrams.html", "flows.html", "play.html", "listen.html",
        "downloads.html", "help.html", "legal.html", "ecosystem.html", "glossary.html", "certs.html", "labs.html",
        "favicon.svg", "manifest.webmanifest", "sw.js", "icon-192.png", "icon-512.png",
    ]
    folders = ["rooms", "assets/css", "assets/js", "guide", "downloads/diagrams", "downloads/sections", "downloads/sheets", "downloads/cards", "downloads/templates"]
    members: list[Path] = []
    for rel in include_roots:
        members.append(ROOT / rel)
    for folder in folders:
        members.extend(path for path in (ROOT / folder).rglob("*") if path.is_file())
    members.append(ROOT / "downloads" / "DISCLAIMER.txt")
    members.append(ROOT / "downloads" / "ai-omni-lab-deck.pptx")
    if destination.exists():
        destination.unlink()
    with zipfile.ZipFile(destination, "w") as archive:
        for path in members:
            if path.resolve() == destination.resolve():
                continue
            if path.suffix == ".zip":
                continue
            info = zipfile.ZipInfo(path.relative_to(ROOT).as_posix())
            info.date_time = (2026, 10, 16, 0, 0, 0)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            archive.writestr(info, path.read_bytes())
    RECORDS.append({"rel": "downloads/ai-omni-lab-offline.zip", "kind": "ZIP", "bytes": destination.stat().st_size, "disclaimer": True})
    print(f"zip {destination.stat().st_size} bytes")


def validate() -> None:
    rooms = 0
    flows = 0
    for _track, room in all_rooms():
        rooms += 1
        if len(room["five"]) != 5 or len(room["quiz"]) != 3 or len(room["actors"]) != 4:
            raise SystemExit("room shape drifted")
        for _q, choices, answer, _why in room["quiz"]:
            if answer < 0 or answer >= len(choices):
                raise SystemExit("quiz answer out of range")
        if room["flow"]:
            flows += 1
    if rooms < 48 or flows < 30:
        raise SystemExit(f"unexpected counts rooms={rooms} flows={flows}")


def main() -> int:
    validate()
    diagrams, flows = build_diagrams()
    for track, room in all_rooms():
        room_page(track, room)
    build_pages(diagrams, flows)
    build_chapter_pages()
    build_play_listen_help()
    build_downloads_assets()
    pages = [{"title": label, "href": href} for _key, label, href, _desc in NAV]
    for track, room in all_rooms():
        pages.append({"title": room["title"], "href": f"rooms/{room_id(track['id'], room['n'])}.html"})
    build_shell_files(pages)
    build_downloads_page()
    build_zip()
    pngs = [record for record in RECORDS if record["kind"] == "PNG" and record["rel"].endswith(".png") and "/diagrams/" in record["rel"]]
    svgs = [record for record in RECORDS if record["kind"] == "SVG" and "/diagrams/" in record["rel"]]
    print(f"rooms={sum(len(track['rooms']) for track in TRACKS)} diagrams={len(svgs)} flows={len(flows)} pngs={len(pngs)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
